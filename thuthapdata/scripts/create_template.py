"""
create_template.py
Tạo file template Excel mẫu cho người dùng nhập liệu từ AI.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

THUTHAPDATADIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_FILE  = os.path.join(THUTHAPDATADIR, "templates", "input_raw_template.xlsx")

def create_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "input_raw"

    # Header
    headers = ["ma_tuyen_sinh", "ten_nganh", "to_hop", "nam", "diem_chuan", "chi_tieu"]
    header_notes = [
        "Mã tuyển sinh trường (VD: BKA, UIT, NEU)",
        "Tên ngành (VD: Công nghệ Thông tin)",
        "Tổ hợp môn (VD: A00, A01, D01)",
        "Năm xét tuyển (VD: 2023, 2024, 2025)",
        "Điểm chuẩn (VD: 27.50)",
        "Chỉ tiêu (VD: 100) — có thể để trống",
    ]

    # Style header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    note_fill   = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    note_font   = Font(color="2F5496", italic=True, name="Calibri", size=10)
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, (h, note) in enumerate(zip(headers, header_notes), start=1):
        # Row 1: tên cột
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.border = border
        cell.alignment = center

        # Row 2: ghi chú
        cell2 = ws.cell(row=2, column=col_idx, value=note)
        cell2.fill   = note_fill
        cell2.font   = note_font
        cell2.border = border
        cell2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Dữ liệu mẫu
    samples = [
        ["BKA",  "Khoa học Máy tính",     "A00", 2025, 27.50, 100],
        ["BKA",  "Khoa học Máy tính",     "A01", 2025, 27.50, 100],
        ["UIT",  "Công nghệ Thông tin",   "A00", 2025, 25.50, 100],
        ["NEU",  "Quản trị Kinh doanh",   "D01", 2025, 27.00, 80],
        ["FTU",  "Ngôn ngữ Anh",          "D01", 2025, 28.25, 60],
    ]

    data_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    data_font = Font(name="Calibri", size=11)

    for row_idx, sample in enumerate(samples, start=3):
        for col_idx, val in enumerate(sample, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill   = data_fill
            cell.font   = data_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Độ rộng cột
    col_widths = [20, 35, 12, 10, 15, 12]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 40

    # Sheet hướng dẫn mã trường
    ws2 = wb.create_sheet("Ma_Tuyen_Sinh")
    ws2.append(["Mã Tuyển Sinh", "Tên Trường", "Khu Vực"])
    from scripts.process import TRUONG_DB
    for ma, info in TRUONG_DB.items():
        ws2.append([ma, info["ten_truong"], info["khu_vuc"]])

    os.makedirs(os.path.dirname(TEMPLATE_FILE), exist_ok=True)
    wb.save(TEMPLATE_FILE)
    print(f"✅ Đã tạo template: {TEMPLATE_FILE}")
    return TEMPLATE_FILE

if __name__ == "__main__":
    create_template()
